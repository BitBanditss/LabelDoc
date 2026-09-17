"""
docx_report.py — Generates an editable DOCX compliance report.
Called by FastAPI; returns the file path of the generated .docx.
"""
from __future__ import annotations
import io, os, base64, tempfile
from datetime import datetime
from typing import Any, Dict

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ── Colour palette ────────────────────────────────────────────────────────────
C_BLUE    = RGBColor(0,   112, 192)
C_DBLUE   = RGBColor(0,   63,  127)
C_GREEN   = RGBColor(22,  163, 74)
C_RED     = RGBColor(220, 38,  38)
C_ORANGE  = RGBColor(234, 88,  12)
C_GRAY    = RGBColor(100, 116, 139)
C_LGRAY   = RGBColor(241, 245, 249)
C_WHITE   = RGBColor(255, 255, 255)
C_DARK    = RGBColor(15,  23,  42)
C_AMBER   = RGBColor(254, 243, 199)


def _set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)


def _set_cell_borders(cell, color='D1D5DB', sz=4):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'),   'single')
        b.set(qn('w:sz'),    str(sz))
        b.set(qn('w:color'), color)
        tcBorders.append(b)
    tcPr.append(tcBorders)


def _para(doc, text, bold=False, size=11, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4):
    p   = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold      = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def _heading(doc, text, level=1, color=C_DBLUE):
    p    = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)
    run  = p.add_run(text)
    run.bold           = True
    run.font.size      = Pt(13 if level == 1 else 11)
    run.font.color.rgb = color
    # bottom border
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '6')
    bot.set(qn('w:color'), '0070C0')
    pBdr.append(bot)
    pPr.append(pBdr)
    return p


def generate_docx(insp: Dict[str, Any]) -> str:
    """Build the DOCX and return the temp file path."""
    doc = Document()

    # ── Page margins ─────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin   = Cm(2.0)
        section.right_margin  = Cm(2.0)

    # ── Default style ─────────────────────────────────────────────────────────
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)

    is_compliant = insp.get('overall_status') == 'COMPLIANT'

    # ── Header ────────────────────────────────────────────────────────────────
    hdr = doc.add_paragraph()
    hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hdr.paragraph_format.space_after = Pt(2)
    r1 = hdr.add_run('LABELDOC')
    r1.bold = True; r1.font.size = Pt(20); r1.font.color.rgb = C_DBLUE
    hdr.add_run('\n')
    r2 = hdr.add_run('Legal Metrology Compliance Inspection Report')
    r2.font.size = Pt(10); r2.font.color.rgb = C_GRAY
    hdr.add_run('\n')
    r3 = hdr.add_run('Government of India · Legal Metrology (Packaged Commodities) Rules, 2011')
    r3.font.size = Pt(9); r3.font.color.rgb = C_GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ── Verdict banner (table with coloured bg) ───────────────────────────────
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    _set_cell_bg(cell, '16A34A' if is_compliant else 'DC2626')
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(8)
    cp.paragraph_format.space_after  = Pt(8)
    icon = '✓' if is_compliant else '✕'
    vr   = cp.add_run(f'{icon}  {"COMPLIANT" if is_compliant else "NON-COMPLIANT"}')
    vr.bold = True; vr.font.size = Pt(16); vr.font.color.rgb = C_WHITE
    cp.add_run(f'\n{insp.get("violation_count", 0)} violation(s) detected  ·  '
               f'Barcode: {"Yes" if insp.get("barcode_calibrated") else "No"}  ·  '
               f'{insp.get("px_per_mm", "—")} px/mm  ·  {insp.get("scanType", "")}').font.color.rgb = C_WHITE

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ── Inspection details grid ───────────────────────────────────────────────
    _heading(doc, 'INSPECTION DETAILS')
    meta = [
        ('Inspection ID',  insp.get('id', '—')),
        ('Timestamp',      datetime.fromisoformat(insp['timestamp']).strftime('%d %b %Y, %I:%M %p')
                           if insp.get('timestamp') else '—'),
        ('Product',        insp.get('productName', '—')),
        ('Scan Type',      insp.get('scanType', '—')),
        ('Inspector',      f"{insp.get('inspectorName','—')} ({insp.get('role','').replace('_',' ')})"),
        ('Jurisdiction',   f"{insp.get('district','—')}, {insp.get('zone','—')}, {insp.get('state','—')}"),
        ('Location',       insp.get('location', '—')),
    ]
    if insp.get('url'):
        meta.append(('URL Scanned', insp['url']))

    tbl2 = doc.add_table(rows=len(meta), cols=2)
    tbl2.style = 'Table Grid'
    tbl2.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (label, value) in enumerate(meta):
        lc = tbl2.cell(i, 0)
        vc = tbl2.cell(i, 1)
        _set_cell_bg(lc, 'F1F5F9')
        lp = lc.paragraphs[0]
        lr = lp.add_run(label)
        lr.bold = True; lr.font.size = Pt(9); lr.font.color.rgb = C_DBLUE
        vp = vc.paragraphs[0]
        vr2 = vp.add_run(str(value))
        vr2.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ── Proof of evidence image ───────────────────────────────────────────────
    image_b64 = insp.get('imageBase64') or ''
    if image_b64:
        _heading(doc, 'PROOF OF EVIDENCE — LABEL PHOTOGRAPH')
        try:
            if ',' in image_b64:
                image_b64 = image_b64.split(',', 1)[1]
            image_b64 = image_b64.strip().replace('\n','').replace('\r','').replace(' ','')
            image_b64 += '=' * (-len(image_b64) % 4)
            img_bytes = base64.b64decode(image_b64)
            img_stream = io.BytesIO(img_bytes)
            doc.add_picture(img_stream, width=Inches(4.5))
            last_para = doc.paragraphs[-1]
            last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(
                f'Captured at time of inspection · '
                f'{datetime.fromisoformat(insp["timestamp"]).strftime("%d %b %Y, %I:%M %p") if insp.get("timestamp") else "—"}'
            )
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size = Pt(8)
            cap.runs[0].font.color.rgb = C_GRAY
        except Exception as e:
            _para(doc, f'[Image could not be embedded: {e}]', color=C_GRAY, size=9)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ── Compliance summary box ────────────────────────────────────────────────
    _heading(doc, 'COMPLIANCE SUMMARY')
    results = insp.get('results', [])
    violations = [(r.get('label') or r.get('field',''), r.get('rule_clause',''), r.get('violations',[]))
                  for r in results if r.get('violations')]

    if not violations:
        p = doc.add_paragraph()
        r = p.add_run('✓  All mandatory declarations comply with Legal Metrology (Packaged Commodities) Rules, 2011')
        r.bold = True; r.font.color.rgb = C_GREEN; r.font.size = Pt(10)
    else:
        sum_tbl = doc.add_table(rows=1, cols=1)
        sum_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        sc = sum_tbl.cell(0, 0)
        _set_cell_bg(sc, 'FEF2F2')
        sp = sc.paragraphs[0]
        sp.paragraph_format.space_before = Pt(4)
        hr = sp.add_run(f'⚠  {insp.get("violation_count", 0)} Violation(s) Detected')
        hr.bold = True; hr.font.color.rgb = C_RED; hr.font.size = Pt(11)
        for (field, clause, viols) in violations:
            np = sc.add_paragraph()
            np.paragraph_format.space_before = Pt(4)
            fr = np.add_run(f'{field}  ({clause})')
            fr.bold = True; fr.font.size = Pt(9); fr.font.color.rgb = RGBColor(127, 29, 29)
            for v in viols:
                vp2 = sc.add_paragraph(style='List Bullet')
                vp2.paragraph_format.left_indent = Cm(0.5)
                vr3 = vp2.add_run(v)
                vr3.font.size = Pt(9); vr3.font.color.rgb = C_RED

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ── Field-by-field audit table ────────────────────────────────────────────
    _heading(doc, 'FIELD-BY-FIELD AUDIT — PCR 2011')
    if results:
        headers = ['Field', 'Value Found on Label', 'Font Size', 'Rule Clause', 'Status']
        widths  = [Cm(4), Cm(7), Cm(2), Cm(3), Cm(3)]
        atbl = doc.add_table(rows=1 + len(results), cols=5)
        atbl.style = 'Table Grid'
        atbl.alignment = WD_TABLE_ALIGNMENT.LEFT

        # Header row
        for j, (h, w) in enumerate(zip(headers, widths)):
            c = atbl.cell(0, j)
            _set_cell_bg(c, '1E3A5F')
            c.width = w
            p2 = c.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(h)
            r2.bold = True; r2.font.size = Pt(9); r2.font.color.rgb = C_WHITE

        # Data rows
        for i, res in enumerate(results):
            row_idx = i + 1
            is_nc   = res.get('status') == 'NON_COMPLIANT'
            row_bg  = 'FFF1F2' if is_nc else 'F0FDF4'

            cells = [
                res.get('label') or res.get('field', '—'),
                str(res.get('value', '—')),
                f"{res['font_height_mm']} mm" if res.get('font_height_mm') is not None else '—',
                res.get('rule_clause', '—'),
                '✕ NON-COMPLIANT' if is_nc else '✓ COMPLIANT',
            ]
            colors = [C_DARK, C_DARK, C_GRAY, C_BLUE,
                      C_RED if is_nc else C_GREEN]
            bolds  = [True, False, False, False, True]

            for j, (val, col, bld) in enumerate(zip(cells, colors, bolds)):
                c = atbl.cell(row_idx, j)
                _set_cell_bg(c, row_bg)
                p3 = c.paragraphs[0]
                r3 = p3.add_run(val)
                r3.font.size = Pt(9)
                r3.font.color.rgb = col
                r3.bold = bld

            # Violations sub-row merged cell
            viols = res.get('violations', [])
            if viols:
                atbl.add_row()
                vrow = atbl.rows[-1]
                merged = vrow.cells[0].merge(vrow.cells[4])
                _set_cell_bg(merged, 'FEE2E2')
                for v in viols:
                    vp3 = merged.add_paragraph()
                    vp3.paragraph_format.left_indent = Cm(0.3)
                    vr4 = vp3.add_run(f'• {v}')
                    vr4.font.size = Pt(8)
                    vr4.font.color.rgb = C_RED

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── Footer disclaimer ─────────────────────────────────────────────────────
    disc = doc.add_paragraph()
    disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dr = disc.add_run(
        '⚠ Prototype compliance assessment. '
        'Final legal determination must be made by the competent Legal Metrology authority.\n'
        f'Generated by LabelDoc · {datetime.now().strftime("%d %b %Y, %I:%M %p")}'
    )
    dr.font.size = Pt(8); dr.font.color.rgb = C_GRAY

    # ── Save to temp file ─────────────────────────────────────────────────────
    tmp = tempfile.NamedTemporaryFile(
        delete=False, suffix='.docx',
        prefix=f'LabelDoc_{insp.get("id","INS")}_'
    )
    doc.save(tmp.name)
    tmp.close()
    return tmp.name